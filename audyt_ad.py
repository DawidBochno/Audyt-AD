#!/usr/bin/env python3
"""Audyt AD - przeglad Active Directory dla administratora. Tylko odczyt
(ADSI/LDAP), nic nie zmienia w domenie.

Zakladki: Podsumowanie (polityka domeny, krbtgt, LAPS, Gosc...), nieaktywne
konta i komputery, konta uprzywilejowane, ryzykowne ustawienia kont
i komputerow (flagi UAC, delegacja, SPN, system bez wsparcia, brak LAPS).
Wynik w oknie i w pliku XLSX.

Uruchomienie: python audyt_ad.py            (GUI)
              python audyt_ad.py --selftest (test logiki)
"""
import datetime
import fnmatch
import os
import re
import sys
import threading

if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))
PLIK_WYKLUCZEN = os.path.join(APP_DIR, "wykluczenia.txt")

FILTR_KONT = "(&(objectCategory=person)(objectClass=user))"
FILTR_KOMPUTEROW = "(objectCategory=computer)"
ATR_KONT = ["sAMAccountName", "displayName", "lastLogonTimestamp", "pwdLastSet",
            "accountExpires", "whenCreated", "userAccountControl", "adminCount",
            "servicePrincipalName", "objectSid", "distinguishedName"]
ATR_KOMPUTEROW = ["name", "operatingSystem", "operatingSystemVersion", "lastLogonTimestamp",
                  "whenCreated", "userAccountControl", "distinguishedName"]
ATR_LAPS = ["msLAPS-PasswordExpirationTime", "ms-Mcs-AdmPwdExpirationTime"]  # nowy, stary LAPS
ATR_DOMENY = ["maxPwdAge", "minPwdLength", "lockoutThreshold", "ms-DS-MachineAccountQuota",
              "msDS-Behavior-Version"]

KOLUMNY = {
    "Podsumowanie": ["Kontrola", "Wynik", "Ocena"],
    "Nieaktywne konta": ["Login", "Imię i nazwisko", "Ostatnie logowanie", "Dni bez logowania",
                         "Utworzone", "Stan", "Jednostka (OU)"],
    "Nieaktywne komputery": ["Nazwa", "System", "Ostatnie logowanie", "Dni bez logowania",
                             "Utworzony", "Stan", "Jednostka (OU)"],
    "Uprzywilejowani": ["Login", "Imię i nazwisko", "Grupy", "Ostatnie logowanie",
                        "Hasło zmienione", "Stan", "Uwagi"],
    "Ryzyka kont": ["Login", "Imię i nazwisko", "Uwagi", "Hasło zmienione",
                    "Ostatnie logowanie", "Jednostka (OU)"],
    "Ryzyka komputerów": ["Nazwa", "System", "Uwagi", "Ostatnie logowanie", "Jednostka (OU)"],
}

UAC_WYLACZONE = 0x2
UAC_KONTROLER = 0x2000  # SERVER_TRUST_ACCOUNT
UAC_DELEGACJA = 0x80000  # TRUSTED_FOR_DELEGATION (nieograniczona)
FLAGI_KONT = [
    (0x20, "hasło niewymagane"),
    (0x80, "hasło z odwracalnym szyfrowaniem"),
    (0x10000, "hasło nigdy nie wygasa"),
    (UAC_DELEGACJA, "delegacja nieograniczona"),
    (0x200000, "tylko szyfrowanie DES"),
    (0x400000, "bez preautoryzacji Kerberos (AS-REP roasting)"),
]
RID_ADMINISTRATOR, RID_GOSC, RID_KRBTGT = 500, 501, 502
POZIOM_DOMENY = {0: "2000", 1: "2003 (przejściowy)", 2: "2003", 3: "2008", 4: "2008 R2",
                 5: "2012", 6: "2012 R2", 7: "2016", 10: "2025"}
FILETIME_1970 = 116444736000000000

# Koniec wsparcia Microsoft (lifecycle). ponytail: tabela reczna - dopisac
# nowe wersje Windows, gdy wyjda; nieznana wersja = brak uwagi.
SYSTEMY = [
    (r"windows xp", "2014-04-08"), (r"windows vista", "2017-04-11"),
    (r"windows 7\b", "2020-01-14"), (r"windows 8\.1", "2023-01-10"), (r"windows 8\b", "2016-01-12"),
    (r"server 2003", "2015-07-14"), (r"server 2008", "2020-01-14"), (r"server 2012", "2023-10-10"),
    (r"server 2016", "2027-01-12"), (r"server 2019", "2029-01-09"),
    (r"server 2022", "2031-10-14"), (r"server 2025", "2034-10-10"),
]
WIN10_LTSC = {10240: "2025-10-14", 14393: "2026-10-13", 17763: "2029-01-09", 19044: "2027-01-12"}
WIN10 = "2025-10-14"  # 22H2, bez rozszerzonych aktualizacji (ESU)
WIN11 = {  # kompilacja: (Home/Pro, Enterprise/Education)
    22000: ("2023-10-10", "2024-10-08"), 22621: ("2024-10-08", "2025-10-14"),
    22631: ("2025-11-11", "2026-11-10"), 26100: ("2026-10-13", "2027-10-12"),
    26200: ("2027-10-12", "2028-10-10"),
}

# ------------------------------------------------------------- konwersje ----


def duza_liczba(v):
    """IADsLargeInteger (HighPart/LowPart) z ADO albo int -> int / None."""
    if v is None:
        return None
    if isinstance(v, int):
        return v
    return (v.HighPart << 32) + (v.LowPart & 0xFFFFFFFF)


def filetime(n):
    """FILETIME (100 ns od 1601, UTC) -> lokalny datetime; 0 / 'nigdy' -> None."""
    if not n or n <= 0 or n >= 0x7FFFFFFFFFFFFFFF:
        return None
    return datetime.datetime.fromtimestamp((n - FILETIME_1970) / 1e7)


def lokalny(v):
    """Data z ADO (UTC, czasem bez strefy) -> lokalny datetime bez strefy."""
    if v is None:
        return None
    if v.tzinfo is None:
        v = v.replace(tzinfo=datetime.timezone.utc)
    return datetime.datetime.fromtimestamp(v.timestamp())


def sid_tekst(b):
    """objectSid (bajty) -> 'S-1-5-21-...-1104'."""
    if not b:
        return ""
    b = bytes(b)
    sub = [int.from_bytes(b[8 + 4 * i:12 + 4 * i], "little") for i in range(b[1])]
    return "S-%d-%d" % (b[0], int.from_bytes(b[2:8], "big")) + "".join("-%d" % s for s in sub)


def jednostka(dn):
    """'CN=Jan,OU=Kadry,OU=Urzad,DC=urzad,DC=local' -> 'Urzad/Kadry'."""
    czesci = re.split(r"(?<!\\),", dn or "")[1:]
    return "/".join(c.split("=", 1)[1] for c in reversed(czesci)
                    if c.upper().startswith(("OU=", "CN=")))


def baza(domena):
    return ",".join("DC=" + d for d in domena.split("."))


def ldap_esc(s):
    """Wartosc do filtra LDAP (RFC 4515)."""
    return "".join("\\%02x" % ord(c) if c in "\\*()\0" else c for c in s)


def pierwszy(v):
    return v[0] if isinstance(v, (tuple, list)) and v else v


def konto(r):
    """Wiersz konta z AD -> rekord."""
    sid = sid_tekst(r.get("objectSid"))
    r = {k: pierwszy(v) for k, v in r.items() if k != "objectSid"}
    uac = r.get("userAccountControl") or 0
    return {
        "nazwa": r.get("sAMAccountName") or "", "opis": r.get("displayName") or "",
        "ostatnie": filetime(duza_liczba(r.get("lastLogonTimestamp"))),
        "haslo": filetime(duza_liczba(r.get("pwdLastSet"))),
        "wygasa": filetime(duza_liczba(r.get("accountExpires"))),
        "utworzone": lokalny(r.get("whenCreated")),
        "uac": uac, "wylaczone": bool(uac & UAC_WYLACZONE),
        "admincount": bool(r.get("adminCount")), "spn": bool(r.get("servicePrincipalName")),
        "rid": int(sid.rsplit("-", 1)[1]) if sid else None,
        "dn": r.get("distinguishedName") or "", "ou": jednostka(r.get("distinguishedName")),
    }


def komputer(r):
    """Wiersz komputera z AD -> rekord."""
    r = {k: pierwszy(v) for k, v in r.items()}
    uac = r.get("userAccountControl") or 0
    return {
        "nazwa": r.get("name") or "", "opis": r.get("operatingSystem") or "",
        "wersja": r.get("operatingSystemVersion") or "",
        "ostatnie": filetime(duza_liczba(r.get("lastLogonTimestamp"))),
        "utworzone": lokalny(r.get("whenCreated")),
        "uac": uac, "wylaczone": bool(uac & UAC_WYLACZONE), "dc": bool(uac & UAC_KONTROLER),
        "laps": max(filter(None, (filetime(duza_liczba(r.get(a))) for a in ATR_LAPS)), default=None),
        "ou": jednostka(r.get("distinguishedName")),
    }

# ---------------------------------------------------------------- AD ----


def pobierz(domena, login="", haslo="", ou=""):
    """Odczyt z AD przez ADSI (ADO). Bez loginu - jako zalogowany uzytkownik
    Windows. Zwraca surowe atrybuty: Konta, Komputery, Grupy, Domena."""
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()  # wywolanie z watku GUI
    conn = win32com.client.Dispatch("ADODB.Connection")
    conn.Provider = "ADsDSOObject"
    if login:
        conn.Properties("User ID").Value = login
        conn.Properties("Password").Value = haslo
        conn.Properties("Encrypt Password").Value = True
    conn.Open("Active Directory Provider")

    def q(base, filtr, atrybuty, zakres="subtree"):
        cmd = win32com.client.Dispatch("ADODB.Command")
        cmd.ActiveConnection = conn
        cmd.Properties("Page Size").Value = 1000  # bez tego max 1000 wynikow
        cmd.CommandText = "<LDAP://%s/%s>;%s;%s;%s" % (domena, base, filtr, ",".join(atrybuty), zakres)
        rs = cmd.Execute()
        rs = rs[0] if isinstance(rs, tuple) else rs  # (Recordset, RecordsAffected)
        out = []
        while not rs.EOF:
            out.append({a: rs.Fields(a).Value for a in atrybuty})
            rs.MoveNext()
        return out

    korzen = baza(domena)
    start = ou or korzen
    komputery = q(start, FILTR_KOMPUTEROW, ATR_KOMPUTEROW)
    for atr in ATR_LAPS:  # atrybut moze nie istniec w schemacie (LAPS niewdrozony)
        try:
            laps = {r["distinguishedName"]: r[atr] for r in
                    q(start, "(&%s(%s=*))" % (FILTR_KOMPUTEROW, atr), ["distinguishedName", atr])}
        except Exception:
            continue
        for k in komputery:
            k[atr] = laps.get(k["distinguishedName"])
    grupy = q(korzen, "(&(objectCategory=group)(adminCount=1))",
              ["name", "objectSid", "distinguishedName"])
    for g in grupy:  # czlonkowie, takze przez grupy zagniezdzone (IN_CHAIN)
        g["czlonkowie"] = [r["distinguishedName"] for r in q(
            korzen, "(&%s(memberOf:1.2.840.113556.1.4.1941:=%s))" % (
                FILTR_KONT, ldap_esc(g["distinguishedName"])), ["distinguishedName"])]
    wynik = {
        "Konta": q(korzen, FILTR_KONT, ATR_KONT),  # cala domena: admini, krbtgt; OU w analizuj()
        "Komputery": komputery,
        "Grupy": grupy,
        "Domena": (q(korzen, "(objectClass=*)", ATR_DOMENY, "base") or [{}])[0],
    }
    conn.Close()
    return wynik

# ------------------------------------------------------------- analiza ----


def nieaktywne(obiekty, dni, teraz, wylaczone=False):
    """Obiekty bez logowania od `dni` dni. Nigdy nie logowane - tylko jesli
    utworzone wczesniej niz `dni` dni temu (nowe konto to nie martwe konto)."""
    granica = teraz - datetime.timedelta(days=dni)
    out = []
    for o in obiekty:
        if o["wylaczone"] and not wylaczone:
            continue
        ost = o["ostatnie"]
        if (ost or o["utworzone"] or teraz) > granica:
            continue
        out.append(dict(o, dni=(teraz - ost).days if ost else None))
    out.sort(key=lambda o: (o["dni"] is not None, -(o["dni"] or 0)))
    return out


def koniec_wsparcia(system, wersja):
    """Data konca wsparcia systemu (date) albo None, gdy nieznana."""
    s = (system or "").lower()
    m = re.search(r"\((\d+)\)", wersja or "")
    build = int(m.group(1)) if m else None
    data = None
    if "windows 11" in s:
        if "ltsc" not in s and build in WIN11:
            data = WIN11[build][any(x in s for x in ("enterprise", "education"))]
    elif "windows 10" in s:
        data = WIN10_LTSC.get(build) if ("ltsc" in s or "ltsb" in s) else WIN10
    else:
        data = next((d for rx, d in SYSTEMY if re.search(rx, s)), None)
    return datetime.date.fromisoformat(data) if data else None


def ryzyka_konta(o, teraz, dni_hasla):
    out = [txt for bit, txt in FLAGI_KONT if o["uac"] & bit]
    if o["spn"]:
        out.append("ma SPN – hasło podatne na Kerberoasting")
    if o["haslo"] and (teraz - o["haslo"]).days > dni_hasla:
        out.append("hasło niezmieniane od %d dni" % (teraz - o["haslo"]).days)
    if o["wygasa"] and o["wygasa"] < teraz:
        out.append("konto wygasło %s" % o["wygasa"].strftime("%Y-%m-%d"))
    return out


def ryzyka_komputera(o, teraz, dni, laps_wdrozony):
    out = []
    koniec = koniec_wsparcia(o["opis"], o["wersja"])
    if koniec and koniec <= teraz.date():
        out.append("system bez wsparcia od %s" % koniec)
    elif koniec and (koniec - teraz.date()).days <= 180:
        out.append("koniec wsparcia systemu %s (za %d dni)" % (koniec, (koniec - teraz.date()).days))
    if not o["dc"]:
        aktywny = o["ostatnie"] and (teraz - o["ostatnie"]).days <= dni
        if laps_wdrozony and not o["laps"]:
            out.append("brak LAPS")
        elif laps_wdrozony and aktywny and o["laps"] < teraz - datetime.timedelta(days=1):
            out.append("LAPS: hasło nieodnowione od %s" % o["laps"].strftime("%Y-%m-%d"))
        if o["uac"] & UAC_DELEGACJA:
            out.append("delegacja nieograniczona")
    return out


def wykluczony(nazwa, wzorce):
    return any(fnmatch.fnmatch(nazwa.lower(), w) for w in wzorce)


def wczytaj_wykluczenia(plik=PLIK_WYKLUCZEN):
    """Jedna nazwa (login / komputer) w linii, dozwolone * i ?, # = komentarz."""
    if not os.path.exists(plik):
        return []
    with open(plik, encoding="utf-8-sig", errors="replace") as f:  # Notatnik ANSI
        return [l.strip().lower() for l in f if l.strip() and not l.lstrip().startswith("#")]


def analizuj(dane, teraz, dni=90, dni_hasla=365, wylaczone=False, wzorce=(), domena="", ou=""):
    """Surowe dane z AD -> {zakladka: wiersze} (kolumny jak w KOLUMNY).
    `ou` zaweza konta w zakladkach nieaktywnych i ryzyk; Uprzywilejowani
    i Podsumowanie zawsze dla calej domeny."""
    konta = [konto(r) for r in dane["Konta"]]
    komputery = [komputer(r) for r in dane["Komputery"]]
    stan = lambda o: "wyłączone" if o["wylaczone"] else "aktywne"
    data = lambda d: d or "nigdy"
    widoczne = lambda o: (wylaczone or not o["wylaczone"]) and not wykluczony(o["nazwa"], wzorce)
    w_ou = [o for o in konta if not ou or o["dn"].lower().endswith(ou.lower())]
    T = {}

    nk = nieaktywne([o for o in w_ou if not wykluczony(o["nazwa"], wzorce)], dni, teraz, wylaczone)
    T["Nieaktywne konta"] = [[o["nazwa"], o["opis"], data(o["ostatnie"]), data(o["dni"]),
                              o["utworzone"], stan(o), o["ou"]] for o in nk]
    nc = nieaktywne([o for o in komputery if not wykluczony(o["nazwa"], wzorce)], dni, teraz, wylaczone)
    T["Nieaktywne komputery"] = [[o["nazwa"], o["opis"], data(o["ostatnie"]), data(o["dni"]),
                                  o["utworzone"], stan(o), o["ou"]] for o in nc]

    # grupy uprzywilejowane = chronione przez AdminSDHolder (adminCount=1)
    grupy_dn = {}
    for g in dane["Grupy"]:
        for dn in g["czlonkowie"]:
            grupy_dn.setdefault(dn, []).append(g["name"])
    T["Uprzywilejowani"] = []
    for o in konta:
        grupy = sorted(set(grupy_dn.get(o["dn"], [])))
        if o["rid"] == RID_KRBTGT or (not grupy and not o["admincount"]):  # krbtgt: w Podsumowaniu
            continue
        uwagi = ryzyka_konta(o, teraz, dni_hasla)
        if not grupy:
            uwagi.insert(0, "adminCount=1 bez członkostwa w grupie uprzywilejowanej (pozostałość)")
        if not o["wylaczone"] and (o["ostatnie"] or o["utworzone"] or teraz) < teraz - datetime.timedelta(days=dni):
            uwagi.insert(0, "nieaktywne od ponad %d dni" % dni)
        T["Uprzywilejowani"].append([o["nazwa"], o["opis"], ", ".join(grupy) or "–",
                                     data(o["ostatnie"]), data(o["haslo"]), stan(o), "; ".join(uwagi)])

    T["Ryzyka kont"] = [[o["nazwa"], o["opis"], "; ".join(u), data(o["haslo"]),
                         data(o["ostatnie"]), o["ou"]]
                        for o in w_ou if widoczne(o) and o["rid"] != RID_KRBTGT
                        for u in [ryzyka_konta(o, teraz, dni_hasla)] if u]
    laps = [o for o in komputery if not o["wylaczone"] and not o["dc"]]
    laps_wdrozony = any(o["laps"] for o in laps)
    T["Ryzyka komputerów"] = [[o["nazwa"], o["opis"], "; ".join(u), data(o["ostatnie"]), o["ou"]]
                              for o in komputery if widoczne(o)
                              for u in [ryzyka_komputera(o, teraz, dni, laps_wdrozony)] if u]

    # ------------------------------------------------ podsumowanie
    d = {k: pierwszy(v) for k, v in dane["Domena"].items()}
    P = [["Domena", domena, "INFO"]]
    poziom = d.get("msDS-Behavior-Version")
    if poziom is not None:
        P.append(["Poziom funkcjonalny domeny", "Windows Server " + POZIOM_DOMENY.get(poziom, str(poziom)),
                  "UWAGA" if poziom < 6 else "OK"])
    if d.get("minPwdLength") is not None:
        P.append(["Minimalna długość hasła (polityka domeny)", "%d znaków" % d["minPwdLength"],
                  "UWAGA" if d["minPwdLength"] < 12 else "OK"])
    if d.get("lockoutThreshold") is not None:
        P.append(["Blokada konta po błędnych hasłach",
                  "po %d próbach" % d["lockoutThreshold"] if d["lockoutThreshold"] else "brak blokady",
                  "OK" if d["lockoutThreshold"] else "UWAGA"])
    wiek = duza_liczba(d.get("maxPwdAge"))
    if wiek is not None:
        P.append(["Maksymalny wiek hasła", "%d dni" % round(-wiek / 864e9)
                  if -2 ** 63 < wiek < 0 else "bez limitu", "INFO"])
    maq = d.get("ms-DS-MachineAccountQuota")
    if maq is not None:
        P.append(["Użytkownicy mogą dodawać komputery do domeny (MachineAccountQuota)",
                  str(maq), "UWAGA" if maq else "OK"])
    po_rid = {o["rid"]: o for o in konta}
    if RID_KRBTGT in po_rid and po_rid[RID_KRBTGT]["haslo"]:
        wiek = (teraz - po_rid[RID_KRBTGT]["haslo"]).days
        P.append(["Hasło konta krbtgt zmienione", "%d dni temu" % wiek, "UWAGA" if wiek > 180 else "OK"])
    if RID_GOSC in po_rid:
        P.append(["Konto Gość", stan(po_rid[RID_GOSC]), "OK" if po_rid[RID_GOSC]["wylaczone"] else "UWAGA"])
    if RID_ADMINISTRATOR in po_rid:
        a = po_rid[RID_ADMINISTRATOR]
        wiek = (teraz - a["haslo"]).days if a["haslo"] else None
        P.append(["Wbudowane konto Administrator (%s)" % a["nazwa"],
                  "%s, hasło zmienione %s" % (stan(a), "%d dni temu" % wiek if wiek is not None else "nigdy"),
                  "UWAGA" if wiek is None or wiek > dni_hasla else "OK"])
    z_laps = sum(1 for o in laps if o["laps"])
    P.append(["LAPS (hasła lokalnych administratorów)",
              "%d z %d komputerów" % (z_laps, len(laps)) if laps_wdrozony else "nie wdrożony",
              "OK" if laps and z_laps == len(laps) else "UWAGA"])
    for zakladka, opis in [("Nieaktywne konta", "Nieaktywne konta (> %d dni)" % dni),
                           ("Nieaktywne komputery", "Nieaktywne komputery (> %d dni)" % dni),
                           ("Uprzywilejowani", "Konta w grupach uprzywilejowanych"),
                           ("Ryzyka kont", "Konta z ryzykownymi ustawieniami"),
                           ("Ryzyka komputerów", "Komputery z uwagami")]:
        n = len(T[zakladka])
        P.append([opis, str(n), "INFO" if zakladka == "Uprzywilejowani" else ("UWAGA" if n else "OK")])
    P.append(["Obiekty w AD", "%d kont, %d komputerów" % (len(konta), len(komputery)), "INFO"])
    if wzorce:
        P.append(["Wykluczenia (wykluczenia.txt)", str(len(wzorce)), "INFO"])
    return dict({"Podsumowanie": P}, **T)


def tekst(v):
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    return "" if v is None else str(v)


def zapisz_xlsx(tabele, sciezka):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    czerwony = PatternFill("solid", fgColor="FFC7CE")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for nazwa, wiersze in tabele.items():
        ws = wb.create_sheet(nazwa)
        ws.append(KOLUMNY[nazwa])
        for c in ws[1]:
            c.font = Font(bold=True)
        for w in wiersze:
            ws.append(w)
            if "UWAGA" in w:
                for c in ws[ws.max_row]:
                    c.fill = czerwony
        for i, kol in enumerate(ws.columns):
            szer = max(len(tekst(c.value)) for c in kol)
            ws.column_dimensions[kol[0].column_letter].width = min(max(szer + 2, 10), 70)
            for c in kol[1:]:
                if isinstance(c.value, datetime.datetime):
                    c.number_format = "yyyy-mm-dd"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(sciezka)


def run(domena, out_dir, log, dni=90, dni_hasla=365, login="", haslo="", wylaczone=False,
        ou="", wzorce=None, zrodlo=None):
    """Odczyt AD, analiza, zapis XLSX. Zwraca {zakladka: wiersze}."""
    if dni < 14:
        log("UWAGA: lastLogonTimestamp jest replikowany z opoznieniem do 14 dni - "
            "wynik ponizej 14 dni jest niedokladny.")
    wzorce = wczytaj_wykluczenia() if wzorce is None else wzorce
    log("Odczyt AD: %s%s ..." % (domena, " / " + ou if ou else ""))
    dane = (zrodlo or pobierz)(domena, login, haslo, ou)
    teraz = datetime.datetime.now()
    tabele = analizuj(dane, teraz, dni, dni_hasla, wylaczone, wzorce, domena, ou)
    for nazwa, wiersze in tabele.items():
        if nazwa != "Podsumowanie":
            log("%s: %d" % (nazwa.replace("ó", "o"), len(wiersze)))
    uwagi = sum(1 for w in tabele["Podsumowanie"] if w[2] == "UWAGA")
    log("Podsumowanie: %d kontroli z UWAGA" % uwagi)
    os.makedirs(out_dir, exist_ok=True)
    plik = os.path.join(out_dir, "audyt_AD_%s_%s.xlsx" % (domena, teraz.strftime("%Y-%m-%d")))
    zapisz_xlsx(tabele, plik)
    log("Zapisano: " + plik)
    return tabele

# ---------------------------------------------------------------- GUI ----


def gui():
    import tkinter as tk
    from tkinter import messagebox, scrolledtext, ttk

    root = tk.Tk()
    root.title("Audyt AD - przeglad Active Directory")
    root.geometry("1150x700")
    pad = dict(padx=6, pady=3)

    v_dom = tk.StringVar(value=os.environ.get("USERDNSDOMAIN", "").lower())
    v_ou = tk.StringVar()
    v_dni = tk.StringVar(value="90")
    v_dni_hasla = tk.StringVar(value="365")
    v_login = tk.StringVar()
    v_haslo = tk.StringVar()
    v_wyl = tk.BooleanVar(value=False)
    v_out = tk.StringVar(value=os.path.join(APP_DIR, "OUTPUT"))
    v_szukaj = tk.StringVar()

    f = ttk.Frame(root)
    f.pack(fill="x", **pad)
    e = lambda var, w, r, c, **kw: ttk.Entry(f, textvariable=var, width=w, **kw).grid(
        row=r, column=c, sticky="w", **pad)
    l = lambda txt, r, c: ttk.Label(f, text=txt).grid(row=r, column=c, sticky="w", **pad)
    l("Domena:", 0, 0); e(v_dom, 28, 0, 1)
    l("OU (opcjonalnie):", 0, 2); e(v_ou, 52, 0, 3)
    l("Login (opcjonalnie):", 1, 0); e(v_login, 28, 1, 1)
    l("Haslo:", 1, 2); e(v_haslo, 28, 1, 3, show="*")
    ttk.Label(f, text="puste = jako zalogowany uzytkownik Windows; inaczej np. URZAD\\admin.  "
              "OU np. OU=Urzad,DC=urzad,DC=local", foreground="gray").grid(
        row=2, column=1, columnspan=3, sticky="w", padx=6)
    g = ttk.Frame(f)
    g.grid(row=3, column=0, columnspan=4, sticky="w")
    ttk.Label(g, text="Bez logowania od (dni):").pack(side="left", **pad)
    ttk.Spinbox(g, textvariable=v_dni, from_=14, to=3650, increment=30, width=6).pack(side="left")
    ttk.Label(g, text="Haslo starsze niz (dni):").pack(side="left", **pad)
    ttk.Spinbox(g, textvariable=v_dni_hasla, from_=30, to=3650, increment=30, width=6).pack(side="left")
    ttk.Checkbutton(g, text="Pokaz tez wylaczone", variable=v_wyl).pack(side="left", **pad)
    ttk.Button(g, text="Wykluczenia...", command=lambda: otworz_wykluczenia()).pack(side="left", **pad)
    l("Folder wyjsciowy:", 4, 0); e(v_out, 60, 4, 1, )

    h = ttk.Frame(root)
    h.pack(fill="x", **pad)
    btn = ttk.Button(h, text="Skanuj")
    btn.pack(side="left", padx=6)
    ttk.Label(h, text="Szukaj:").pack(side="left", padx=(30, 4))
    ttk.Entry(h, textvariable=v_szukaj, width=30).pack(side="left")

    nb = ttk.Notebook(root)
    nb.pack(fill="both", expand=True, **pad)
    drzewa, dane = {}, {}
    szer = {"Uwagi": 420, "Grupy": 230, "Jednostka (OU)": 220, "Imię i nazwisko": 170,
            "Kontrola": 420, "Wynik": 300, "System": 190}

    def sortuj(t, kol, odwrotnie):
        klucz = lambda s: (0, int(s)) if s.isdigit() else (1, s.lower())
        el = sorted(t.get_children(), key=lambda k: klucz(t.set(k, kol)), reverse=odwrotnie)
        for i, k in enumerate(el):
            t.move(k, "", i)
        t.heading(kol, command=lambda: sortuj(t, kol, not odwrotnie))

    for nazwa, kolumny in KOLUMNY.items():
        ramka = ttk.Frame(nb)
        nb.add(ramka, text=nazwa)
        t = ttk.Treeview(ramka, columns=kolumny, show="headings")
        for k in kolumny:
            t.heading(k, text=k, command=lambda t=t, k=k: sortuj(t, k, False))
            t.column(k, width=szer.get(k, 115), anchor="w")
        t.tag_configure("uwaga", background="#FFD7D7")
        t.bind("<Double-1>", lambda ev, t=t, kol=kolumny: t.focus() and messagebox.showinfo(
            "Szczegoly", "\n".join("%s: %s" % kv for kv in zip(kol, t.item(t.focus(), "values")))))
        sb = ttk.Scrollbar(ramka, orient="vertical", command=t.yview)
        sx = ttk.Scrollbar(ramka, orient="horizontal", command=t.xview)
        t.configure(yscrollcommand=sb.set, xscrollcommand=sx.set)
        sx.pack(side="bottom", fill="x")
        sb.pack(side="right", fill="y")
        t.pack(side="left", fill="both", expand=True)
        drzewa[nazwa] = t

    log_box = scrolledtext.ScrolledText(root, height=5)
    log_box.pack(fill="x", **pad)

    def log(msg):
        def put():
            log_box.insert("end", str(msg) + "\n")
            log_box.see("end")
        root.after(0, put)

    def wypelnij(*_):
        fraza = v_szukaj.get().strip().lower()
        for nazwa, t in drzewa.items():
            t.delete(*t.get_children())
            for w in dane.get(nazwa, []):
                wart = [tekst(x) for x in w]
                if fraza and not any(fraza in x.lower() for x in wart):
                    continue
                t.insert("", "end", values=wart, tags=("uwaga",) if "UWAGA" in wart else ())
            if nazwa != "Podsumowanie" and nazwa in dane:
                nb.tab(t.master, text="%s (%d)" % (nazwa, len(dane[nazwa])))

    v_szukaj.trace_add("write", wypelnij)

    def otworz_wykluczenia():
        if not os.path.exists(PLIK_WYKLUCZEN):
            with open(PLIK_WYKLUCZEN, "w", encoding="utf-8") as fw:
                fw.write("# Konta i komputery pomijane w audycie (poza zakladka Uprzywilejowani).\n"
                         "# Jedna nazwa w linii, mozna uzyc * i ?, np.:\n# svc_*\n# SKANER-01\n")
        os.startfile(PLIK_WYKLUCZEN)

    def start():
        log_box.delete("1.0", "end")
        dom, ou = v_dom.get().strip(), v_ou.get().strip()
        if not dom or "." not in dom:
            return log("Podaj nazwe domeny, np. urzad.local")
        if ou and "DC=" not in ou.upper():
            return log("OU podaj jako pelna sciezke, np. OU=Urzad,DC=urzad,DC=local")
        try:
            dni, dni_hasla = int(v_dni.get()), int(v_dni_hasla.get())
        except ValueError:
            return log("Liczby dni musza byc liczbami calkowitymi.")
        btn.config(state="disabled")

        def work():
            try:
                wynik = run(dom, v_out.get().strip('" '), log, dni, dni_hasla,
                            v_login.get().strip(), v_haslo.get(), v_wyl.get(), ou)
                def pokaz():  # w watku GUI - wypelnij() czyta `dane` przy wyszukiwaniu
                    dane.clear()
                    dane.update(wynik)
                    wypelnij()
                root.after(0, pokaz)
            except PermissionError as e:
                log("BLAD: nie mozna zapisac %s - zamknij ten plik w Excelu." % e.filename)
            except Exception as e:  # com_error, brak sieci, bledne haslo
                info = getattr(e, "excepinfo", None)
                e = info[2].strip() if info and info[2] else e
                log("BLAD: nie udalo sie odczytac AD (%s). Czy komputer widzi kontroler "
                    "domeny, a login i haslo sa poprawne?" % (e,))
            finally:
                root.after(0, lambda: btn.config(state="normal"))

        threading.Thread(target=work, daemon=True).start()

    btn.config(command=start)
    if "--selftest" in sys.argv:
        root.after(200, root.destroy)
    import aktualizacja
    aktualizacja.start(root, "DawidBochno/Audyt-AD", "main", "audyt_ad.py")
    root.mainloop()

# ------------------------------------------------------------- selftest ----


def selftest():
    import tempfile
    import openpyxl

    class LI:  # IADsLargeInteger z ADO
        def __init__(self, n):
            self.HighPart, self.LowPart = n >> 32, n & 0xFFFFFFFF
            if self.LowPart >= 2 ** 31:  # COM oddaje LowPart ze znakiem
                self.LowPart -= 2 ** 32

    def ft(dt):
        return int(dt.timestamp() * 1e7) + FILETIME_1970

    def sid(*sub):
        return bytes([1, len(sub), 0, 0, 0, 0, 0, 5]) + b"".join(s.to_bytes(4, "little") for s in sub)

    teraz = datetime.datetime.now().replace(microsecond=0)
    d = lambda n: teraz - datetime.timedelta(days=n)
    utc = lambda n: d(n).astimezone(datetime.timezone.utc)

    # konwersje
    assert duza_liczba(LI(ft(d(200)))) == ft(d(200)) and filetime(ft(d(200))) == d(200)
    assert duza_liczba(LI(-36288000000000)) == -36288000000000  # maxPwdAge 42 dni
    assert filetime(0) is None and filetime(0x7FFFFFFFFFFFFFFF) is None
    assert sid_tekst(sid(32, 544)) == "S-1-5-32-544"
    assert sid_tekst(memoryview(sid(21, 1, 2, 3, 1104))) == "S-1-5-21-1-2-3-1104"
    assert jednostka(r"CN=Kowalski\, Jan,OU=Kadry,OU=Urzad,DC=urzad,DC=local") == "Urzad/Kadry"
    assert jednostka("CN=PC01,CN=Computers,DC=urzad,DC=local") == "Computers"
    assert baza("urzad.gmina.local") == "DC=urzad,DC=gmina,DC=local"
    assert ldap_esc(r"CN=Admins (IT)\, Urzad*") == r"CN=Admins \28IT\29\5c, Urzad\2a"
    assert lokalny(datetime.datetime(2026, 1, 1, 12)) == lokalny(
        datetime.datetime(2026, 1, 1, 12, tzinfo=datetime.timezone.utc))
    assert wykluczony("SVC_Backup", ["svc_*"]) and not wykluczony("jnowak", ["svc_*"])

    # koniec wsparcia
    kw = koniec_wsparcia
    assert kw("Windows 7 Professional", "6.1 (7601)") == datetime.date(2020, 1, 14)
    assert kw("Windows 8.1 Pro", "6.3 (9600)") == datetime.date(2023, 1, 10)
    assert kw("Windows Server 2012 R2 Standard", "6.3 (9600)") == datetime.date(2023, 10, 10)
    assert kw("Windows 10 Pro", "10.0 (19045)") == datetime.date(2025, 10, 14)
    assert kw("Windows 10 Enterprise LTSC", "10.0 (17763)") == datetime.date(2029, 1, 9)
    assert kw("Windows 11 Pro", "10.0 (22631)") == datetime.date(2025, 11, 11)
    assert kw("Windows 11 Enterprise", "10.0 (22631)") == datetime.date(2026, 11, 10)
    assert kw("Windows 11 Pro", "10.0 (29999)") is None and kw("", "") is None
    assert kw("Windows Server 2022 Standard", "10.0 (20348)") == datetime.date(2031, 10, 14)

    # sztuczna domena
    D = "DC=urzad,DC=local"

    def u(nazwa, ost, utw=900, uac=512, haslo=10, rid=1100, spn=None, ac=None, wygasa=None, opis=None):
        return {"sAMAccountName": nazwa, "displayName": opis or nazwa.title(),
                "lastLogonTimestamp": LI(ft(d(ost))) if ost is not None else None,
                "pwdLastSet": LI(ft(d(haslo))) if haslo is not None else LI(0),
                "accountExpires": LI(ft(d(wygasa))) if wygasa is not None else LI(0x7FFFFFFFFFFFFFFF),
                "whenCreated": utc(utw), "userAccountControl": uac, "adminCount": ac,
                "servicePrincipalName": spn, "objectSid": sid(21, 1, 2, 3, rid),
                "distinguishedName": "CN=%s,OU=Kadry,%s" % (nazwa, D)}

    def k(nazwa, ost, system="Windows 11 Pro", wersja="10.0 (29999)", uac=4096, laps=30):
        return {"name": nazwa, "operatingSystem": system, "operatingSystemVersion": wersja,
                "lastLogonTimestamp": LI(ft(d(ost))), "whenCreated": utc(900),
                "userAccountControl": uac, "distinguishedName": "CN=%s,OU=PC,%s" % (nazwa, D),
                "msLAPS-PasswordExpirationTime": LI(ft(d(-laps))) if laps is not None else None}

    dane = {
        "Konta": [
            u("aktywny", 3), u("stary", 200), u("nigdy", None, utw=400), u("nowy", None, utw=5),
            u("wylaczony", 300, uac=514), u("granica", 91),
            u("Administrator", 2, haslo=800, rid=500, uac=512 | 0x10000, ac=1),
            u("Gosc", None, uac=514 | 0x20, rid=501, haslo=None), u("krbtgt", None, uac=514, rid=502,
                                                                   haslo=1000, spn=("kadmin/changepw",), ac=1),
            u("svc_sql", 1, uac=512 | 0x10000, spn=("MSSQLSvc/sql01",), haslo=900),
            u("asrep", 1, uac=512 | 0x400000), u("odwracalne", 1, uac=512 | 0x80),
            u("bylyadmin", 1, ac=1), u("adm_it", 120, ac=1), u("wygasle", 1, wygasa=10),
        ],
        "Komputery": [
            k("PC-OK", 1), k("PC-W10", 2, "Windows 10 Pro", "10.0 (19045)"),
            k("PC-NOLAPS", 1, laps=None), k("PC-LAPS-STARE", 1, laps=-20),
            k("PC-MARTWY", 200, laps=-150), k("DC01", 1, "Windows Server 2025 Standard",
                                               "10.0 (26100)", uac=0x2000 | UAC_DELEGACJA, laps=None),
            k("SRV-DELEG", 1, "Windows Server 2025 Standard", "10.0 (26100)", uac=4096 | UAC_DELEGACJA),
            k("PC-WYL", 500, "Windows 7 Professional", "6.1 (7601)", uac=4098, laps=None),
        ],
        "Grupy": [
            {"name": "Administratorzy domeny", "objectSid": sid(21, 1, 2, 3, 512),
             "distinguishedName": "CN=Administratorzy domeny,CN=Users," + D,
             "czlonkowie": ["CN=Administrator,OU=Kadry," + D, "CN=adm_it,OU=Kadry," + D]},
            {"name": "Administratorzy", "objectSid": sid(32, 544),
             "distinguishedName": "CN=Administratorzy,CN=Builtin," + D,
             "czlonkowie": ["CN=Administrator,OU=Kadry," + D, "CN=adm_it,OU=Kadry," + D]},
        ],
        "Domena": {"maxPwdAge": LI(-36288000000000), "minPwdLength": 7, "lockoutThreshold": 0,
                   "ms-DS-MachineAccountQuota": 10, "msDS-Behavior-Version": 7},
    }
    T = analizuj(dane, teraz, 90, 365, False, ["wygasl*"], "urzad.local")
    assert list(T) == list(KOLUMNY), list(T)
    assert all(len(w) == len(KOLUMNY[n]) for n, ws in T.items() for w in ws)
    assert [w[0] for w in T["Nieaktywne konta"]] == ["nigdy", "stary", "adm_it", "granica"], T["Nieaktywne konta"]
    assert T["Nieaktywne konta"][0][2:4] == ["nigdy", "nigdy"] and T["Nieaktywne konta"][1][3] == 200
    assert [w[0] for w in T["Nieaktywne komputery"]] == ["PC-MARTWY"]

    upr = {w[0]: w for w in T["Uprzywilejowani"]}
    assert sorted(upr) == ["Administrator", "adm_it", "bylyadmin"], sorted(upr)
    assert upr["adm_it"][2] == "Administratorzy, Administratorzy domeny"
    assert upr["adm_it"][6].startswith("nieaktywne od ponad 90 dni")
    assert upr["bylyadmin"][6].startswith("adminCount=1 bez") and upr["bylyadmin"][2] == "–"
    assert "hasło nigdy nie wygasa" in upr["Administrator"][6]

    rk = {w[0]: w[2] for w in T["Ryzyka kont"]}
    assert "krbtgt" not in rk and "Gosc" not in rk and "wygasle" not in rk, rk
    assert "ma SPN" in rk["svc_sql"] and "hasło niezmieniane od 900 dni" in rk["svc_sql"]
    assert rk["asrep"] == "bez preautoryzacji Kerberos (AS-REP roasting)"
    assert rk["odwracalne"] == "hasło z odwracalnym szyfrowaniem"
    assert "aktywny" not in rk
    T2 = analizuj(dane, teraz, 90, 365, True, [], "urzad.local")
    rk2 = {w[0]: w[2] for w in T2["Ryzyka kont"]}
    assert rk2["wygasle"].startswith("konto wygasło") and "hasło niewymagane" in rk2["Gosc"]

    rc = {w[0]: w[2] for w in T["Ryzyka komputerów"]}
    assert "PC-OK" not in rc and "DC01" not in rc and "PC-WYL" not in rc, rc
    assert rc["PC-W10"] == "system bez wsparcia od 2025-10-14"
    assert rc["PC-NOLAPS"] == "brak LAPS"
    assert rc["PC-LAPS-STARE"].startswith("LAPS: hasło nieodnowione od")
    assert rc["SRV-DELEG"] == "delegacja nieograniczona"
    assert "PC-MARTWY" not in rc  # nieaktywny - stary LAPS to nie blad LAPS

    P = {w[0]: w[1:] for w in T["Podsumowanie"]}
    assert P["Minimalna długość hasła (polityka domeny)"] == ["7 znaków", "UWAGA"]
    assert P["Blokada konta po błędnych hasłach"] == ["brak blokady", "UWAGA"]
    assert P["Maksymalny wiek hasła"] == ["42 dni", "INFO"]
    assert P["Poziom funkcjonalny domeny"] == ["Windows Server 2016", "OK"]
    assert P["Hasło konta krbtgt zmienione"] == ["1000 dni temu", "UWAGA"]
    assert P["Konto Gość"] == ["wyłączone", "OK"]
    assert P["Wbudowane konto Administrator (Administrator)"][1] == "UWAGA"
    assert P["LAPS (hasła lokalnych administratorów)"] == ["5 z 6 komputerów", "UWAGA"], P
    assert P["Użytkownicy mogą dodawać komputery do domeny (MachineAccountQuota)"] == ["10", "UWAGA"]
    assert P["Wykluczenia (wykluczenia.txt)"] == ["1", "INFO"]

    # OU zaweza nieaktywne i ryzyka, ale nie administratorow i Podsumowania
    T4 = analizuj(dane, teraz, 90, 365, False, [], "urzad.local", ou="ou=PUSTA," + D)
    assert not T4["Nieaktywne konta"] and not T4["Ryzyka kont"]
    assert len(T4["Uprzywilejowani"]) == 3 and "Hasło konta krbtgt zmienione" in {w[0] for w in T4["Podsumowanie"]}
    T5 = analizuj(dane, teraz, 90, 365, False, [], "urzad.local", ou="OU=Kadry," + D)
    assert len(T5["Nieaktywne konta"]) == 4

    # bez LAPS w domenie: nie zasypuje kazdego komputera uwaga "brak LAPS"
    bez = dict(dane, Komputery=[dict(x, **{"msLAPS-PasswordExpirationTime": None}) for x in dane["Komputery"]])
    T3 = analizuj(bez, teraz, 90, 365, False, [], "urzad.local")
    assert not any("LAPS" in w[2] for w in T3["Ryzyka komputerów"])
    assert {w[0]: w[1] for w in T3["Podsumowanie"]}["LAPS (hasła lokalnych administratorów)"] == "nie wdrożony"

    # wykluczenia z pliku, run + XLSX
    tmp = tempfile.mkdtemp()
    with open(os.path.join(tmp, "w.txt"), "w", encoding="utf-8") as fw:
        fw.write("# komentarz\n\n  STARY \nsvc_*\n")
    assert wczytaj_wykluczenia(os.path.join(tmp, "w.txt")) == ["stary", "svc_*"]
    assert wczytaj_wykluczenia(os.path.join(tmp, "brak.txt")) == []
    with open(os.path.join(tmp, "ansi.txt"), "w", encoding="cp1250") as fw:
        fw.write("ksiegowa\nżółw\n")
    assert wczytaj_wykluczenia(os.path.join(tmp, "ansi.txt"))[0] == "ksiegowa"
    lines = []
    wyn = run("urzad.local", tmp, lines.append, dni=7, wzorce=["stary"],
              zrodlo=lambda dom, l, h, ou: dane)
    assert all(ord(ch) < 128 for line in lines for ch in line), lines
    assert any("UWAGA" in x for x in lines) and any("Ryzyka komputerow:" in x for x in lines), lines
    assert "stary" not in [w[0] for w in wyn["Nieaktywne konta"]]
    plik = [f for f in os.listdir(tmp) if f.endswith(".xlsx")]
    assert plik == ["audyt_AD_urzad.local_%s.xlsx" % teraz.strftime("%Y-%m-%d")], plik
    wb = openpyxl.load_workbook(os.path.join(tmp, plik[0]))
    assert wb.sheetnames == list(KOLUMNY)
    ws = wb["Nieaktywne konta"]
    assert [c.value for c in ws[1]] == KOLUMNY["Nieaktywne konta"]
    assert ws["A2"].value == "nigdy" and ws["C2"].value == "nigdy"
    assert isinstance(ws["E2"].value, datetime.datetime) and ws["E2"].number_format == "yyyy-mm-dd"
    ws = wb["Podsumowanie"]
    uw = [r for r in ws.iter_rows(min_row=2) if r[2].value == "UWAGA"]
    assert uw and all(c.fill.fgColor.rgb.endswith("FFC7CE") for c in uw[0])

    import aktualizacja
    aktualizacja.selftest()
    print("selftest OK")


if __name__ == "__main__":
    if "--selftest" in sys.argv:
        selftest()
        if "--gui" in sys.argv:
            gui()
    else:
        gui()
